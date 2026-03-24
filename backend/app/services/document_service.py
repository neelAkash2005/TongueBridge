import re
from pathlib import Path
from uuid import uuid4
import logging

from fastapi import HTTPException, UploadFile

from app.services.translation_service import translate_text

# Setup logging
logger = logging.getLogger(__name__)


UPLOAD_DIR = Path("temp_docs")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".doc", ".docx", ".ppt", ".pptx"}


CONTRACTION_MAP = {
    "can't": "cannot",
    "won't": "will not",
    "don't": "do not",
    "didn't": "did not",
    "isn't": "is not",
    "aren't": "are not",
    "i'm": "I am",
    "it's": "it is",
    "that's": "that is",
    "there's": "there is",
    "we're": "we are",
    "you're": "you are",
    "they're": "they are",
    "i've": "I have",
    "you've": "you have",
}


def _expand_contractions(text: str) -> str:
    result = text
    for short_form, long_form in CONTRACTION_MAP.items():
        pattern = re.compile(f"\\b{re.escape(short_form)}\\b", re.IGNORECASE)
        def replace_func(match):
            if match.group(0)[0].isupper():
                return long_form[0].upper() + long_form[1:]
            return long_form
        result = pattern.sub(replace_func, result)
    return result


def _to_sentence_case(text: str) -> str:
    if not text:
        return text
    return text[0].upper() + text[1:]


def _apply_tone_preserving_text(text: str, tone: str) -> str:
    raw = text.strip()
    if not raw:
        return ""

    if tone == "formal":
        expanded = _expand_contractions(raw)
        punctuated = expanded if re.search(r"[.!?]$", expanded) else f"{expanded}."
        return f"Kindly note: {_to_sentence_case(punctuated)}"

    if tone == "casual":
        soft = re.sub(r"\bdo not\b", "don't", raw, flags=re.IGNORECASE)
        soft = re.sub(r"\bcannot\b", "can't", soft, flags=re.IGNORECASE)
        punctuated = soft if re.search(r"[.!?]$", soft) else f"{soft}!"
        return f"Hey! {_to_sentence_case(punctuated)}"

    return _to_sentence_case(raw)


def _split_text_for_translation(text: str, max_chunk_size: int = 250) -> list[str]:
    normalized_text = re.sub(r"\s+", " ", text).strip()
    if not normalized_text:
        return []

    parts = re.split(r"(?<=[.!?])\s+", normalized_text)
    chunks: list[str] = []
    current_chunk = ""

    for part in parts:
        if not current_chunk:
            current_chunk = part
            continue

        candidate = f"{current_chunk} {part}".strip()
        if len(candidate) <= max_chunk_size:
            current_chunk = candidate
        else:
            chunks.append(current_chunk)
            current_chunk = part

    if current_chunk:
        chunks.append(current_chunk)

    return chunks


def _extract_text_from_txt(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="ignore").strip()


def _extract_text_from_pdf(file_path: Path) -> str:
    from pypdf import PdfReader
    
    reader = PdfReader(str(file_path))
    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n".join(page_texts).strip()


def _extract_text_from_docx(file_path: Path) -> str:
    from docx import Document
    
    document = Document(str(file_path))
    paragraph_text = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraph_text).strip()


def _extract_text_from_doc(file_path: Path) -> str:
    import olefile
    
    if not olefile.isOleFile(str(file_path)):
        raise ValueError("Invalid .doc file format.")

    extracted_segments: list[str] = []
    with olefile.OleFileIO(str(file_path)) as ole:
        stream_names = ["WordDocument", "1Table", "0Table"]
        for stream_name in stream_names:
            if not ole.exists(stream_name):
                continue

            stream_data = ole.openstream(stream_name).read()

            utf16_matches = re.findall(rb"(?:[\x20-\x7E]\x00){4,}", stream_data)
            ascii_matches = re.findall(rb"[\x20-\x7E]{6,}", stream_data)

            for match in utf16_matches:
                decoded = match.decode("utf-16le", errors="ignore").strip()
                if decoded:
                    extracted_segments.append(decoded)

            for match in ascii_matches:
                decoded = match.decode("latin-1", errors="ignore").strip()
                if decoded:
                    extracted_segments.append(decoded)

    cleaned_text = "\n".join(extracted_segments)
    cleaned_text = re.sub(r"\s+", " ", cleaned_text).strip()
    return cleaned_text


def _extract_text_from_pptx(file_path: Path) -> str:
    from pptx import Presentation
    
    presentation = Presentation(str(file_path))
    slide_texts: list[str] = []

    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())

    return "\n".join(slide_texts).strip()


async def translate_document(
    file: UploadFile,
    source_lang: str,
    target_lang: str,
    tone: str = "neutral",
) -> dict:
    original_name = file.filename or "document.txt"
    suffix = Path(original_name).suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Only .txt, .pdf, .doc, .docx, .ppt, and .pptx files are supported.")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    upload_filename = f"upload_{uuid4().hex}{suffix}"
    upload_path = UPLOAD_DIR / upload_filename

    try:
        upload_path.write_bytes(file_bytes)
        if suffix == ".txt":
            extracted_text = _extract_text_from_txt(upload_path)
        elif suffix == ".pdf":
            extracted_text = _extract_text_from_pdf(upload_path)
        elif suffix == ".docx":
            extracted_text = _extract_text_from_docx(upload_path)
        elif suffix == ".doc":
            extracted_text = _extract_text_from_doc(upload_path)
        else:
            extracted_text = _extract_text_from_pptx(upload_path)
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Failed to read document content: {error}") from error
    finally:
        try:
            upload_path.unlink()
        except Exception:
            pass

    if not extracted_text:
        raise HTTPException(status_code=400, detail="No readable text found in the uploaded document.")

    # Apply tone preserving if not neutral
    text_to_translate = extracted_text
    if tone != "neutral":
        text_to_translate = _apply_tone_preserving_text(extracted_text, tone)

    text_chunks = _split_text_for_translation(text_to_translate)
    if not text_chunks:
        raise HTTPException(status_code=400, detail="Document does not contain translatable text.")

    try:
        translated_chunks = []
        for chunk in text_chunks:
            translated_chunks.append(translate_text(chunk, source_lang, target_lang))

        translated_text = " ".join(chunk for chunk in translated_chunks if chunk).strip()
    except Exception as e:
        logger.error(f"Translation failed: {e}")
        raise HTTPException(status_code=500, detail=f"Translation failed: {str(e)}")

    if not translated_text:
        raise HTTPException(status_code=500, detail="Translation failed for document content.")

    translated_filename = f"translated_{uuid4().hex}.txt"
    translated_path = UPLOAD_DIR / translated_filename
    translated_path.write_text(translated_text, encoding="utf-8")

    return {
        "filename": translated_filename,
        "extracted_text": extracted_text,
        "translated_text": translated_text,
        "download_url": f"/download/{translated_filename}",
    }
